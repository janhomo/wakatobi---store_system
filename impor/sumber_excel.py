from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

N, S, R, K0 = "Normal", "Servis", "Rusak", "Belum diisi"
G, KP, KPQ = "Gudang", "Karpenter", "Karpenter?"
# (lembar, no_kertas, nama, merek, model, unit, kondisi, lokasi, perlu_dicek, catatan, unit_diisi)
rows = []
def add(l, no, nama, merek, model, unit, kond, lok=G, cek=0, cat="", filled=0):
    rows.append((l, no, nama, merek, model, unit, kond, lok, cek, cat, filled))

# ---- Lembar 1 Makita impak (hal. 1) ----
L = "H1 Makita"
add(L,1,"Bor impak 18V","Makita","DTW 300 RTJ",1,N)
add(L,2,"Bor impak 18V","Makita","DTW 300 RTJ",2,N)
add(L,3,"Bor impak 18V","Makita","DTW 700 RTJ",3,N)
add(L,4,"Bor impak 18V","Makita","DTW 700 RTJ",4,N)
add(L,5,"Bor magnet 1050W","Makita","HB 350-2QAA",1,N,cek=1,cat="Model ragu terbaca.")
add(L,6,"Bor duduk","Proline","DF 4113",2,N,cek=1,cat="Model ragu. Hanya 1 baris dan bernomor 2: apakah unit 1 belum tercatat?")
add(L,7,"Bor cas 1.5Ah","Makita","DF 333D",1,R)
add(L,8,"Bor cas 1.5Ah","Makita","DF 333D",2,N)
add(L,9,"Bor cas 1.5Ah","Makita","DF 333D",3,N)
add(L,10,"Bor cas 1.5Ah","Makita","DF 333D",4,S)
add(L,11,"Bor cas 1.5Ah","Makita","DF 333D",5,R,cek=1,cat="Model tertulis 'DF 33 D', disamakan jadi DF 333D.")
add(L,12,"Bor cas 18V 5.0Ah","Makita","DDF 484",1,N,cek=1,cat="No. unit terbaca 0 atau 1. Dipilih 1.")
add(L,13,"Bor cas 18V 5.0Ah","Makita","DDF 484",2,N)
add(L,14,"Bor cas 7.2V 1.5Ah","Makita","DF 012 D",1,K0,cek=1,cat="Kondisi tidak dicentang.")
# ---- Halaman 2 ----
L = "H2"
add(L,15,"Rebar tier","Makita","DTR 180 Z",1,K0,cek=1,filled=1,cat="No. unit dan kondisi kosong. Unit diisi 1 oleh saya. Nama ada kata buram (kuning?).")
add(L,16,"Rebar tier","Makita","DTR 180 Z",2,K0,cek=1,filled=1,cat="No. unit dan kondisi kosong. Unit diisi 2 oleh saya.")
add(L,17,"Rebar tier kuning","Machine code","",1,R,cek=1,cat="Merek tertulis 'Machine code' (dari plat, atau pengganti TIDAK ADA?). Model terhapus.")
add(L,18,"Rebar tier kuning","Machine code","",2,N,cek=1,cat="Sama dengan baris 17.")
add(L,19,"Rebar tier kuning","Machine code","",3,N,cek=1,cat="Sama dengan baris 17.")
add(L,20,"Rebar tier kuning","Machine code","",4,N,cek=1,cat="Sama dengan baris 17.")
add(L,21,"Rebar tier kuning","Machine code","",5,N,cek=1,cat="Sama dengan baris 17.")
add(L,22,"Blower","Krisbow","SRVP 300",3,N)
add(L,23,"Blower","Krisbow","SRVP 300",4,N)
add(L,24,"Blower","Makita","UB 1103",1,N)
add(L,25,"Blower","Nankai","",5,N,cek=1,cat="Model kosong.")
add(L,26,"Blower","Nankai","",2,N,cek=1,cat="Model kosong.")
add(L,27,"Jiksaw","Bosch","GST 680",2,N)
add(L,28,"Jiksaw","Bosch","GST 680",1,R)
# ---- Halaman 3 ----
L = "H3"
add(L,29,"Jiksaw","Makita","4351 FCT",3,N,lok=KP,cek=1,filled=1,cat="No. unit kosong. Diisi 3 oleh saya. Tanda K di tepi.")
add(L,30,"Hot gun","Bosch","GHG 2063",1,N)
add(L,31,"Hot gun","Bosch","GHG 2063",2,N)
add(L,32,"Hot gun","Bosch","GHG 2063",3,N)
add(L,33,"Mixer","Komax","Machine code",3,N,cek=1,cat="Model tertulis 'Machine code' (dari plat, atau pengganti TIDAK ADA?).")
add(L,34,"Mixer","Makita","UT 2204",4,N)
add(L,35,"Mixer","Wipro","W 5980",1,N)
add(L,36,"Mixer","Wipro","W 5980",2,N)
add(L,37,"Skap","Makita","KP 0800",2,N,lok=KP,cat="Tanda K di tepi.")
add(L,38,"Skap","Makita","KP 0800",3,N,lok=KP,cek=1,cat="Tanda K di tepi. No. unit terbaca 3 atau 5. Dipilih 3.")
add(L,39,"Skap","Makita","KP 0800",1,N)
add(L,40,"Somel","Makita","N 5900 B",1,S,cek=1,cat="Servis dicentang, kotak Rusak juga tampak bertanda. Cek mesinnya.")
add(L,41,"Somel","Makita","N 5900 B",2,N)
add(L,42,"Gergaji mesin","Bosch","GSA 120",1,N)
# ---- Halaman 4 ----
L = "H4"
add(L,43,"Amplas dinding 18V","Makita","PSL 800",1,N,cek=1,cat="Model ragu (PSL atau DSL 800).")
add(L,44,"Amplas tembok","Wipro","W 4225",2,N,cek=1,cat="Hanya 1 baris dan bernomor 2: apakah unit 1 belum tercatat?")
add(L,45,"Las PVC","Weldy","HT 1600",1,N)
add(L,46,"Somel cas 18V","Makita","DHS 660",1,N)
add(L,47,"Somel cas 18V","Makita","DHS 660",2,N)
add(L,48,"Somel cas 18V","Makita","DCS 5532",3,N,cek=1,cat="Model ragu (DCS 553 Z?).")
add(L,49,"Laser","Bosch","GLL 3-60 XG",1,N)
add(L,50,"Laser","Bosch","GLL 3-60 XG",2,N)
add(L,51,"Laser","Leica","0810",3,S,cek=1,cat="Kotak Normal tercoret, Servis dicentang. Model ragu.")
add(L,52,"Laser","Tactix","25411",4,N)
add(L,53,"Potong ubin","Rubi","TC 125",1,N)
add(L,54,"Core drill 2400W","Wipro","WP 02-205",1,N)
add(L,55,"Belt sander 440W","Makita","M 9400",2,N,cek=1,cat="Hanya 1 baris dan bernomor 2. Model ragu.")
add(L,56,"Bor breaker kecil","Bosch","GBH 2-28 DFV",1,N,cek=1,cat="Di kertas tertulis 'Bor Breker' tanpa 'kecil'. Disamakan dengan baris 57-58 (model sama).")
# ---- Halaman 5 ----
L = "H5"
add(L,57,"Bor breaker kecil","Bosch","GBH 2-28 DFV",2,N,lok=KPQ,cek=1,cat="Ada goresan di tepi foto, mungkin K yang terpotong.")
add(L,58,"Bor breaker kecil","Bosch","GBH 2-28 DFV",3,N,lok=KPQ,cek=1,cat="Goresan di tepi foto, mungkin K terpotong. No. unit 3 kembar dengan Lembar Bosch baris 2.")
add(L,59,"Bor breaker sedang","Bosch","GBH 5-40 D",4,R)
add(L,60,"Bor breaker sedang","Bosch","GBH 5-40 D",5,R)
add(L,61,"Breker besar","Makita","HM 1810",2,S)
add(L,62,"Breker besar","Makita","HM 1810",3,N)
add(L,63,"Breker besar","Makita","HM 1810",4,S)
add(L,64,"Breker besar","Makita","HM 1810",1,S)
add(L,65,"Breker besar","Makita","HM 1810",7,S,cek=1,cat="No. unit terbaca 7 atau 1. Dipilih 7 karena 1 sudah dipakai.")
add(L,66,"Breker besar","Makita","HM 1810",9,N,cek=1,cat="Model ditimpa tulisan. Servis dicoret, Normal dicentang.")
add(L,67,"Breker besar","Makita","HM 1812",5,S)
add(L,68,"Breker besar","Makita","HM 1812",10,N)
add(L,69,"Breker besar","Makita","HM 1306",6,K0,cek=1,cat="Kondisi tidak dicentang.")
add(L,70,"Breker besar","Bosch","GSH 27 VC",8,K0,cek=1,cat="Kondisi tidak dicentang. Model ragu.")
# ---- Lembar Bosch (hal. 1) ----
L = "H1 Bosch"
add(L,1,"Bor cas 18V 2.0Ah","Bosch","GSB 185 LI",1,N)
add(L,2,"Bor breaker kecil","Bosch","GBH 2-28 DFV",3,N,lok=KPQ,cek=1,cat="Goresan di tepi foto, mungkin K terpotong. No. unit 3 kembar dengan H5 baris 58: mesin yang sama?")
add(L,3,"Bor listrik 350W","Modern","M.2100 C",2,N)
add(L,4,"Bor listrik 350W","Modern","M.2100 C",3,N)
add(L,5,"Bor listrik 350W","Makita","TW 0350",9,N,cek=1,cat="No. unit terbaca 9 atau 3. Dipilih 9 karena 3 sudah dipakai Modern.")
add(L,6,"Bor listrik 750W","Bosch","GSB 16 RE",1,N)
add(L,7,"Bor listrik 750W","Bosch","GSB 16 RE",7,R)
add(L,8,"Bor listrik 750W","Bosch","GSB 16 RE",4,N)
add(L,9,"Bor listrik 750W","Bosch","GSB 16 RE",8,N)
add(L,10,"Bor listrik 400W","Bosch","GBM 400",5,N,cek=1,cat="Daya tertimpa tulisan (400 W).")
add(L,11,"Bor listrik 400W","Bosch","GBM 400",6,N)
add(L,12,"Potong keramik 860W","Makita","4100 NB",1,N)
add(L,13,"Potong besi 2000W","Makita","2414 NB",1,N)
add(L,14,"Pembengkok besi","Kowloon","GF 25",1,N)
# ---- Lembar Gurinda (hal. 1) ----
L = "H1 Gurinda"
add(L,1,"Pemotong multifungsi","Bosch","GOP 30-28",1,N)
add(L,2,"Gurinda besar","Bosch","GWS 2200",1,R)
add(L,3,"Gurinda besar","Bosch","GWS 2200",2,N)
add(L,4,"Gurinda besar","Bosch","GWS 2200",3,N,cek=1,cat="Model tertimpa tulisan (2800 menjadi 2200).")
add(L,5,"Gurinda besar","Bosch","GWS 2200",4,N)
add(L,6,"Vakum","Bosch","GAS 12-25",1,N)
add(L,7,"Vakum","Bosch","GAS 12-25",3,N,lok=KP,cat="Tanda K di tepi.")
add(L,8,"Vakum","Krisbow","KW 1800306",2,N)
add(L,9,"Vakum/Blower","Makita","UB 1103",1,N,cek=1,cat="Mungkin sama dengan H2 baris 24 (Blower Makita UB 1103 no. 1).")
add(L,10,"Las 160A","Lakoni","Falcon 160 e",1,N)
add(L,11,"Las 200A","Lakoni","HAW TIG 200 e",2,N,cek=1,cat="Model ragu. Hanya 1 baris dan bernomor 2.")
add(L,12,"Air compressor","Imola","IMOLA-225",3,N)
add(L,13,"Air compressor","Izumi","OD 1024 1 HP",1,N,lok=KP,cek=1,cat="Tanda K di tepi. Model OD 1024, sedangkan baris 14 tertulis DD 1024: yang mana benar?")
add(L,14,"Air compressor","Izumi","DD 1024",2,N,cek=1,cat="Lihat baris 13: OD atau DD?")

assert len(rows) == 98, len(rows)

wb = Workbook()
F = "Arial"
hdr_fill = PatternFill("solid", start_color="1B2226")
thin = Side(style="thin", color="C9C8C0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
yellow = PatternFill("solid", start_color="FFF2A8")
red = PatternFill("solid", start_color="F6C9C5")
blue = PatternFill("solid", start_color="D6E6F2")

# ---------- Sheet Mesin ----------
ws = wb.active
ws.title = "Mesin"
heads = ["No", "Lembar / Hal.", "No. kertas", "Nama mesin", "Merek", "Model", "No. unit",
         "Kondisi (kertas)", "Lokasi", "Perlu dicek", "Catatan", "Jenis (rumus)", "Cek kembar (rumus)"]
widths = [5, 13, 9, 24, 14, 18, 8, 15, 12, 10, 62, 44, 14]
for c, (h, w) in enumerate(zip(heads, widths), 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 30
last = len(rows) + 1
for i, r in enumerate(rows, 2):
    l, no, nama, merek, model, unit, kond, lok, cek, cat, filled = r
    vals = [i - 1, l, no, nama, merek, model if model else None, unit, kond, lok,
            "Ya" if cek else None, cat if cat else None,
            f'=D{i}&" | "&E{i}&" | "&F{i}',
            f'=IF(COUNTIFS($D$2:$D${last},D{i},$G$2:$G${last},G{i})>1,"KEMBAR","")']
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font = Font(name=F, size=10)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (4, 11, 12)),
                                   horizontal="center" if c in (1, 3, 7, 8, 9, 10, 13) else "left")
    if filled:
        ws.cell(row=i, column=7).fill = yellow
    if kond == K0:
        ws.cell(row=i, column=8).fill = red
    if lok != G:
        ws.cell(row=i, column=9).fill = blue
    if cek:
        ws.cell(row=i, column=10).fill = yellow
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:M{last}"

# ---------- Sheet Ringkasan ----------
rs = wb.create_sheet("Ringkasan")
rs.column_dimensions["A"].width = 44
rs.column_dimensions["B"].width = 12
rs.column_dimensions["C"].width = 60
items = [
    ("Jumlah baris di kertas", f"=COUNTA(Mesin!D2:D{last})", "Dari 7 foto (hal. 1-5 dan dua lembar hal. 1 lain)."),
    ("Jumlah jenis (nama + merek + model)", f"=SUMPRODUCT(1/COUNTIF(Mesin!L2:L{last},Mesin!L2:L{last}))", "Salah ketik model memecah satu jenis jadi dua."),
    ("Kondisi: Normal", f'=COUNTIF(Mesin!H2:H{last},"Normal")', ""),
    ("Kondisi: Servis", f'=COUNTIF(Mesin!H2:H{last},"Servis")', ""),
    ("Kondisi: Rusak", f'=COUNTIF(Mesin!H2:H{last},"Rusak")', ""),
    ("Kondisi: Belum diisi", f'=COUNTIF(Mesin!H2:H{last},"Belum diisi")', "Harus diisi saat label ditempel."),
    ("Cek jumlah kondisi = jumlah baris", '=IF(SUM(B6:B9)=B4,"OK","SALAH")', ""),
    ("Lokasi: Karpenter (tanda K jelas)", f'=COUNTIF(Mesin!I2:I{last},"Karpenter")', "K di luar tabel = mesin di karpenter (kata Anda)."),
    ("Lokasi: Karpenter? (K terpotong foto)", f'=COUNTIF(Mesin!I2:I{last},"Karpenter?")', "Perlu dipastikan."),
    ("Lokasi: Gudang (tanpa tanda K)", f'=COUNTIF(Mesin!I2:I{last},"Gudang")', "Asumsi: tanpa K berarti di gudang."),
    ("Baris perlu dicek", f'=COUNTIF(Mesin!J2:J{last},"Ya")', "Filter kolom J = Ya."),
    ("Baris No. unit kembar dalam nama yang sama", f'=COUNTIF(Mesin!M2:M{last},"KEMBAR")', "Dihitung per baris, jadi 2 baris kembar = 2."),
    ("No. unit yang saya isi sendiri (kosong di kertas)", f'=SUMPRODUCT(--ISNUMBER(SEARCH("diisi",Mesin!K2:K{last})))', "Sel kuning di kolom G."),
]
rs["A1"] = "Ringkasan daftar mesin awal"
rs["A1"].font = Font(name=F, bold=True, size=13)
rs["A3"], rs["B3"], rs["C3"] = "Hal", "Nilai", "Catatan"
for c in ("A3", "B3", "C3"):
    rs[c].font = Font(name=F, bold=True, color="FFFFFF", size=10)
    rs[c].fill = hdr_fill
for j, (a, b, c) in enumerate(items, 4):
    rs.cell(row=j, column=1, value=a).font = Font(name=F, size=10)
    cb = rs.cell(row=j, column=2, value=b); cb.font = Font(name=F, size=10, bold=True)
    cb.alignment = Alignment(horizontal="center")
    rs.cell(row=j, column=3, value=c).font = Font(name=F, size=10)
    for k in (1, 2, 3):
        rs.cell(row=j, column=k).border = border

# ---------- Sheet Baca saya ----------
bs = wb.create_sheet("Baca saya")
bs.column_dimensions["A"].width = 120
notes = [
    ("Daftar mesin awal dari Lembar Catat Mesin Gudang", True),
    ("Sumber: 7 foto lembar yang diisi penjaga gudang. Tanggal dan nama pencatat kosong di semua lembar.", False),
    ("", False),
    ("Aturan yang dipakai", True),
    ("1. No. unit = angka yang ditulis penjaga di kolom 'No. Seri' (atas keputusan Anda). Bukan seri pabrik.", False),
    ("2. Tanda K di luar tabel = mesin berada di Karpenter. 'Karpenter?' = goresan terpotong tepi foto, mungkin K.", False),
    ("3. Mesin tanpa tanda K dianggap di Gudang. Ini asumsi saya.", False),
    ("4. Ejaan nama mesin saya samakan dan huruf besar kecilnya saya rapikan. Merek dan model ditulis seperti di kertas.", False),
    ("5. Kondisi di kertas belum boleh dipercaya penuh. Isi ulang dari mesin yang dipegang saat label ditempel.", False),
    ("", False),
    ("Warna sel", True),
    ("Kuning di kolom 'No. unit' = nomor kosong di kertas, diisi oleh saya.", False),
    ("Merah muda di kolom 'Kondisi' = kondisi tidak dicentang.", False),
    ("Biru di kolom 'Lokasi' = di Karpenter (atau mungkin).", False),
    ("Kuning di kolom 'Perlu dicek' = ada keraguan baca atau data janggal. Lihat kolom Catatan.", False),
    ("", False),
    ("Cara pakai", True),
    ("Koreksi langsung di sheet 'Mesin'. Kolom 'Jenis' dan 'Cek kembar' dihitung otomatis dari nama, merek, model, dan no. unit.", False),
    ("Filter kolom 'Perlu dicek' = Ya untuk melihat baris yang harus dicek. Sheet 'Ringkasan' ikut berubah otomatis.", False),
]
for i, (t, b) in enumerate(notes, 1):
    c = bs.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=b, size=12 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.move_sheet("Baca saya", offset=-2)
wb.active = 1
out = "/tmp/claude-0/-home-claude/eda91aa1-27da-5de8-9165-aabef528fe76/scratchpad/data/Daftar_Mesin_Awal.xlsx"
wb.save(out)
print("saved", out)
